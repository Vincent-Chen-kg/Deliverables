from typing import Any
from docx import Document
from docx.table import Table

from .utils import add_header_with_superscript_date


class CostReportDocumentBuilder:
    """Renders cost tables and constructs docx documents based on templates."""

    @staticmethod
    def build_cost_table(
        doc: Document,
        sorted_records: list[dict[str, Any]],
        total_cumulative: float,
        total_current: float,
        cum_start_date: str,
        start_date: str,
        end_date: str
    ) -> Table:
        """Renders the cost table inside the reference Word document."""
        table = doc.tables[1]

        # Column Headers
        hdr_cells = table.rows[0].cells
        add_header_with_superscript_date(
            hdr_cells[1].paragraphs[0], "Cumulative Linked Account Total", cum_start_date, end_date
        )
        add_header_with_superscript_date(
            hdr_cells[2].paragraphs[0], "Latest Total", start_date, end_date
        )

        # Totals Row
        total_cells = table.rows[1].cells
        total_cells[1].text = f"${total_cumulative:,.2f}"
        total_cells[2].text = f"${total_current:,.2f}"

        # Account Data Rows
        for index, record in enumerate(sorted_records):
            target_index = index + 2
            if target_index < len(table.rows):
                row = table.rows[target_index]
            else:
                row = table.add_row()

            row_cells = row.cells
            acc_name = record.get('account_name')
            acc_id = record['account_id']

            row_cells[0].text = f"{acc_name} ({acc_id})" if acc_name and acc_name != "Unknown / External" else acc_id
            row_cells[1].text = f"${record['cumulative']:,.2f}"
            row_cells[2].text = f"${record['current']:,.2f}"

        return table
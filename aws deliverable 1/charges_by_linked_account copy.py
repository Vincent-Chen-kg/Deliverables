import boto3
from pathlib import Path
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from botocore.exceptions import ClientError
import json

# Net of credits, refunds, and negotiated discounts, i.e. what actually gets invoiced.

TEMPLATE_DIR =  Path(__file__).resolve().parents[1]/"templates"
COST_METRIC = "NetUnblendedCost"
file_path= "exceptions.json"
cum_start_date = "2026-06-13"

try:
    with open(file_path, "r", encoding="utf-8") as file:
        exceptions = json.load(file)["exceptions"]
except FileNotFoundError:
    print(f"File '{file_path}' does not exist.")
    exceptions = {} 

def get_account_names(account_ids: set) -> dict[str, str]:
    """Fetches AWS account names using AWS Organizations API."""
    org_client = boto3.client('organizations')
    account_names = {}
 
    for acc_id in account_ids:
        try:
            res = org_client.describe_account(AccountId=acc_id)
            account_names[acc_id] = res['Account']['Name']
        except ClientError:
            account_names[acc_id] = "Unknown / External"
 
    return account_names
 
def get_date_parts(date_str: str) -> tuple[str, int, str, int]:
    """Returns (month_name, day_num, suffix, year) for superscript formatting."""
    dt = datetime.strptime(date_str, "%Y-%m-%d")
    day = dt.day
 
    if 11 <= day <= 13:
        suffix = 'th'
    else:
        suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(day % 10, 'th')
 
    return dt.strftime('%B'), day, suffix, dt.year
 
def add_header_with_superscript_date(paragraph, prefix: str, start_date_str: str, end_date_str: str):
    """Appends header text with ordinal day suffixes formatted as superscripts."""
    p1_month, p1_day, p1_suf, p1_year = get_date_parts(start_date_str)
    p2_month, p2_day, p2_suf, p2_year = get_date_parts(end_date_str)
 
    paragraph.add_run(f"{prefix} (")
    paragraph.add_run(f"{p1_month} {p1_day}")
    r_s1 = paragraph.add_run(p1_suf)
    r_s1.font.superscript = True
    paragraph.add_run(f", {p1_year} to ")

    paragraph.add_run(f"{p2_month} {p2_day}")
    r_s2 = paragraph.add_run(p2_suf)
    r_s2.font.superscript = True
    paragraph.add_run(f", {p2_year})")

    for run in paragraph.runs:
        run.bold = True

def generate_styled_aws_cost_docx(start_date: str, end_date: str, query_end_date: str, title: str ):

    print(f"start date: {start_date}")
    print(f"end date: {end_date}")
    print(f"cumulative start date:{cum_start_date}")

    # Cost Explorer treats TimePeriod.End as exclusive, so query one day past the
    # last day being reported.

    ce_client = boto3.client('ce')

    # 1. Query Cumulative Cost
    response_cum = ce_client.get_cost_and_usage(
        TimePeriod={'Start': cum_start_date, 'End': query_end_date},
        Granularity='MONTHLY',
        Metrics=[COST_METRIC],
        GroupBy=[{'Type': 'DIMENSION', 'Key': 'LINKED_ACCOUNT'}]
    )

    # 2. Query Current Period Cost
    response_curr = ce_client.get_cost_and_usage(
        TimePeriod={'Start': start_date, 'End': query_end_date},
        Granularity='MONTHLY',
        Metrics=[COST_METRIC],
        GroupBy=[{'Type': 'DIMENSION', 'Key': 'LINKED_ACCOUNT'}]
    )

    # 3. Aggregate Data
    account_data = {}

    for result in response_cum.get('ResultsByTime', []):
        for group in result.get('Groups', []):
            acc_id = group['Keys'][0]
            if acc_id in exceptions:
                continue
            amount = float(group['Metrics'][COST_METRIC]['Amount'])
            account_data[acc_id] = account_data.get(acc_id, {'cumulative': 0.0, 'current': 0.0})
            account_data[acc_id]['cumulative'] += amount

    for result in response_curr.get('ResultsByTime', []):
        for group in result.get('Groups', []):
            acc_id = group['Keys'][0]
            if acc_id in exceptions:
                continue
            amount = float(group['Metrics'][COST_METRIC]['Amount'])
            if acc_id not in account_data:
                account_data[acc_id] = {'cumulative': 0.0, 'current': 0.0}
            account_data[acc_id]['current'] += amount

    # Fetch Account Names
    account_names = get_account_names(set(account_data.keys()))

    # Totals & Sorting
    total_cumulative = sum(item['cumulative'] for item in account_data.values())
    total_current = sum(item['current'] for item in account_data.values())

    sorted_records = sorted(
        [{'account_id': k, 'account_name': account_names.get(k, ''), **v} for k, v in account_data.items()],
        key=lambda x: x['cumulative'],
        reverse=True
    )

    # 4. Build Document
    doc = Document(TEMPLATE_DIR / "template1.docx")

    build_cost_table(
        doc,
        sorted_records,
        total_cumulative,
        total_current,
        cum_start_date,
        start_date,
        end_date,
    )

    doc.save(title)
    
    print(f"Report successfully saved to {title}")

def build_cost_table(doc, sorted_records, total_cumulative, total_current,
                     cum_start_date: str, start_date: str, end_date: str):
    """Renders the cost table using the reference document's table formatting."""
    table = doc.tables[1]

    header_row = table.rows[0]
    hdr_cells = header_row.cells

    # Column Headers
    add_header_with_superscript_date(hdr_cells[1].paragraphs[0], "Cumulative Linked Account Total", cum_start_date, end_date)
    add_header_with_superscript_date(hdr_cells[2].paragraphs[0], "Latest Total", start_date, end_date)


    # 5. Total Costs Row
    total_row = table.rows[1]
    total_cells = total_row.cells

    total_cells[1].text = f"${total_cumulative:,.2f}"
    total_cells[2].text = f"${total_current:,.2f}"

    # 6. Populate Data Rows
    for index, record in enumerate(sorted_records):
        if index+2 <len(table.rows):
            row =  table.rows[index+2] 
        else:
            table.add_row()
            print("An extra row was created")
            
        row_cells = row.cells

        if record['account_name'] and record['account_name'] != "Unknown / External":
            row_cells[0].text = f"{record['account_name']} ({record['account_id']})"
        else:
            row_cells[0].text = record['account_id']

        row_cells[1].text = f"${record['cumulative']:,.2f}"
        row_cells[2].text = f"${record['current']:,.2f}"

    return table

def get_previous_month_parameters(date):
    first_day_this_month = date.replace(day=1)
    last_day_prev_month = first_day_this_month - timedelta(days=1)
    first_day_prev_month = last_day_prev_month.replace(day=1)

    #days in YYYY-MM-DD format
    start_date = first_day_prev_month.strftime('%Y-%m-%d')
    end_date = last_day_prev_month.strftime('%Y-%m-%d')
    query_end_date = first_day_this_month.strftime('%Y-%m-%d')
    return {"start_date":start_date, "end_date": end_date,"query_end_date": query_end_date, "title": f"{last_day_prev_month.strftime("%Y_%B")}_AWS_Cost_Report.docx"}

if __name__ == '__main__':

    today = datetime.now().date()
    query = get_previous_month_parameters(today)
    print(query)

    generate_styled_aws_cost_docx(**query)